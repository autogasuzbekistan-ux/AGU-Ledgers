"""Click hisobot faylini o'qish: har kontragentning kunlik jami click summasi."""
from openpyxl import load_workbook


def parse_click_file(path):
    """Qaytaradi: ({kontragent_nomi: jami_som}, kurs)"""
    wb = load_workbook(path, data_only=True)
    ws = wb.active

    headers = {c: str(ws.cell(row=1, column=c).value).strip()
               for c in range(1, ws.max_column + 1) if ws.cell(row=1, column=c).value}

    kurs = None
    for row in ws.iter_rows():
        for cell in row:
            if isinstance(cell.value, str) and 'курс' in cell.value.lower():
                kc = ws.cell(row=cell.row, column=cell.column + 1)
                if isinstance(kc.value, (int, float)):
                    kurs = kc.value

    # "Jami" qatori qattiq kodlanmagan - dasturiy aniqlanadi: qaysi qator
    # qiymati o'zidan yuqoridagi hammasi yig'indisiga teng bo'lsa, o'sha jami.
    totals_row = None
    for r in range(2, ws.max_row + 1):
        checked = matches = 0
        for col in headers:
            val = ws.cell(row=r, column=col).value
            if val is None:
                continue
            checked += 1
            col_sum = sum((ws.cell(row=rr, column=col).value or 0) for rr in range(2, r))
            if isinstance(val, (int, float)) and abs(val - col_sum) < 1:
                matches += 1
        if checked >= 3 and matches == checked:
            totals_row = r
            break

    if totals_row is None:
        raise ValueError(f"'{path}': jami qatori topilmadi - fayl tuzilishi kutilganidan farq qiladi")

    return {headers[c]: (ws.cell(row=totals_row, column=c).value or 0) for c in headers}, kurs
