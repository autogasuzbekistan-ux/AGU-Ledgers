"""Kunlik hisobot faylini o'qish: naqd so'm / naqd dollar / plastik / click,
har kuni keladigan kontragentlar bo'yicha.

MUHIM: bu faylda 'Qog'oz' qatori = naqd DOLLAR (real fayllarda tasdiqlangan,
kassa daftari bilan solishtirib). 'naqt' = naqd so'm.
"""
from openpyxl import load_workbook

ROW_LABELS = {
    'naqt': 'naqd_som',
    "qog'oz": 'naqd_dollar',
    'plastik': 'terminal',
    'click': 'click',
}


def parse_hisobot_file(path):
    """Qaytaradi: {kontragent_nomi: {naqd_som, naqd_dollar, terminal, click}}

    Maydon nomi 'terminal' (fayldagi 'Plastik' qatoridan) - report_parser.py
    va ledger.DailyEntry bilan bir xil nomlanish, ikkalasi ham
    on_reupload_confirm'da bir xil kalit ('terminal') orqali DailyEntry'ga
    yoziladi."""
    wb = load_workbook(path, data_only=True)
    ws = wb.active

    names = {c: str(ws.cell(row=1, column=c).value).strip()
             for c in range(2, ws.max_column + 1) if ws.cell(row=1, column=c).value}

    rows = {}
    for r in range(2, ws.max_row + 1):
        label = ws.cell(row=r, column=1).value
        if label:
            key = str(label).strip().lower()
            if key in ROW_LABELS:
                rows[ROW_LABELS[key]] = r

    result = {}
    for c, name in names.items():
        result[name] = {
            field: (ws.cell(row=row_num, column=c).value or 0)
            for field, row_num in rows.items()
        }
    return result
