"""Professional ko'rinishdagi hisobot fayllarini (xlsx) shakllantirish.

Telegram chat xabari uzun ro'yxatlar uchun noqulay (qisqartirish kerak,
o'qish qiyin) - shuning uchun "Qarzdorlar" kabi to'liq ro'yxatlar
formatlangan Excel fayl sifatida yuboriladi (bot.py orqali).
"""
from datetime import date as _date

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

HEADER_FILL = PatternFill("solid", fgColor="4472C4")
HEADER_FONT = Font(bold=True, color="FFFFFF")
USD_FORMAT = '$#,##0.00'
SOM_FORMAT = '#,##0'


def _style_header_row(ws, row, headers):
    for col, text in enumerate(headers, start=1):
        cell = ws.cell(row=row, column=col, value=text)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal="center")


def _set_column_widths(ws, widths):
    for col, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(col)].width = width


def _write_debtors_sheet(ws, debtors, sana):
    ws.title = "Qarzdorlar"

    ws["A1"] = f"Qarzdorlar ro'yxati - {sana.isoformat()}"
    ws["A1"].font = Font(bold=True, size=14)
    ws.merge_cells("A1:E1")
    ws["A2"] = f"Jami qarzdorlar: {len(debtors)} ta"
    ws["A2"].font = Font(italic=True, color="666666")

    header_row = 4
    headers = ["№", "Kontragent", "Oxirgi to'lov sanasi", "Oxirgi to'lov summasi ($)", "Qarz ($)"]
    _style_header_row(ws, header_row, headers)

    total_qarz = 0
    r = header_row
    for i, (nomi, qarz, oxirgi_tolov_sana, oxirgi_tolov_summa) in enumerate(debtors, start=1):
        r = header_row + i
        ws.cell(row=r, column=1, value=i)
        ws.cell(row=r, column=2, value=nomi)
        ws.cell(row=r, column=3, value=oxirgi_tolov_sana.isoformat() if oxirgi_tolov_sana else "ma'lum emas")
        oxirgi_cell = ws.cell(row=r, column=4, value=oxirgi_tolov_summa)
        oxirgi_cell.number_format = USD_FORMAT
        qarz_cell = ws.cell(row=r, column=5, value=qarz)
        qarz_cell.number_format = USD_FORMAT
        total_qarz += qarz

    total_row = r + 1
    ws.cell(row=total_row, column=2, value="JAMI").font = Font(bold=True)
    total_cell = ws.cell(row=total_row, column=5, value=total_qarz)
    total_cell.font = Font(bold=True)
    total_cell.number_format = USD_FORMAT

    _set_column_widths(ws, [6, 32, 18, 22, 16])
    ws.freeze_panes = f"A{header_row + 1}"


def _write_daily_payments_sheet(ws, daily_rows):
    ws.title = "Kunlik to'lovlar"

    ws["A1"] = "Kunlik jami tushum (barcha kontragentlar bo'yicha)"
    ws["A1"].font = Font(bold=True, size=14)
    ws.merge_cells("A1:F1")

    header_row = 3
    headers = ["Sana", "Naqd so'm", "Click", "Terminal", "Naqd dollar", "Jami ($)"]
    _style_header_row(ws, header_row, headers)

    total = {"naqd_som": 0, "click": 0, "terminal": 0, "naqd_dollar": 0, "usd": 0}
    r = header_row
    for i, (sana, naqd_som, click, terminal, naqd_dollar, jami_usd) in enumerate(daily_rows, start=1):
        r = header_row + i
        ws.cell(row=r, column=1, value=sana.isoformat())
        ws.cell(row=r, column=2, value=naqd_som).number_format = SOM_FORMAT
        ws.cell(row=r, column=3, value=click).number_format = SOM_FORMAT
        ws.cell(row=r, column=4, value=terminal).number_format = SOM_FORMAT
        ws.cell(row=r, column=5, value=naqd_dollar).number_format = USD_FORMAT
        ws.cell(row=r, column=6, value=jami_usd).number_format = USD_FORMAT
        total["naqd_som"] += naqd_som
        total["click"] += click
        total["terminal"] += terminal
        total["naqd_dollar"] += naqd_dollar
        total["usd"] += jami_usd

    total_row = r + 1
    ws.cell(row=total_row, column=1, value="JAMI").font = Font(bold=True)
    for col, key, fmt in [(2, "naqd_som", SOM_FORMAT), (3, "click", SOM_FORMAT),
                           (4, "terminal", SOM_FORMAT), (5, "naqd_dollar", USD_FORMAT),
                           (6, "usd", USD_FORMAT)]:
        cell = ws.cell(row=total_row, column=col, value=total[key])
        cell.font = Font(bold=True)
        cell.number_format = fmt

    _set_column_widths(ws, [14, 16, 16, 16, 16, 16])
    ws.freeze_panes = f"A{header_row + 1}"


def build_debtors_report(debtors, daily_rows, path, sana=None):
    """Ikkita varaqli professional Excel hisobot yaratadi:

    - "Qarzdorlar": debtors (bot_logic.top_debtors natijasi -
      [(kontragent_nomi, qarz_dollar, oxirgi_tolov_sana_yoki_None,
      oxirgi_tolov_summa_dollar), ...]) - eng katta qarzdan boshlab.
    - "Kunlik to'lovlar": daily_rows (bot_logic.daily_payments_summary
      natijasi) - barcha kontragentlar bo'yicha kun kesimida jami tushum.

    `path`ga yozadi va shu yo'lni qaytaradi."""
    sana = sana or _date.today()
    wb = Workbook()
    _write_debtors_sheet(wb.active, debtors, sana)
    _write_daily_payments_sheet(wb.create_sheet(), daily_rows)
    wb.save(path)
    return path
