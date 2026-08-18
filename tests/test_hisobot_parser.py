from openpyxl import Workbook

from parsers import parse_hisobot_file


def _build_hisobot_file(path):
    wb = Workbook()
    ws = wb.active
    ws.append([None, "Rashid aka", "Alisher aka"])
    ws.append(["naqt", 100000, 50000])
    ws.append(["Qog'oz", 20, 10])
    ws.append(["Plastik", 30000, 0])
    ws.append(["click", 15000, 25000])
    wb.save(path)


def test_parse_hisobot_file_maps_row_labels_per_kontragent(tmp_path):
    path = tmp_path / "hisobot_01_08_2026.xlsx"
    _build_hisobot_file(path)

    result = parse_hisobot_file(path)

    assert result == {
        "Rashid aka": {
            "naqd_som": 100000,
            "naqd_dollar": 20,
            "terminal": 30000,
            "click": 15000,
        },
        "Alisher aka": {
            "naqd_som": 50000,
            "naqd_dollar": 10,
            "terminal": 0,
            "click": 25000,
        },
    }


def test_parse_hisobot_file_row_labels_are_case_insensitive(tmp_path):
    wb = Workbook()
    ws = wb.active
    ws.append([None, "Rashid aka"])
    ws.append(["NAQT", 100000])
    ws.append(["qog'oz", 20])
    ws.append(["PLASTIK", 0])
    ws.append(["Click", 5000])
    path = tmp_path / "hisobot_case.xlsx"
    wb.save(path)

    result = parse_hisobot_file(path)

    assert result["Rashid aka"] == {
        "naqd_som": 100000,
        "naqd_dollar": 20,
        "terminal": 0,
        "click": 5000,
    }


def test_parse_hisobot_file_unknown_row_is_added_to_naqd_dollar(tmp_path):
    # Real faylda uchragan holat: standart 4 qatordan tashqari "Vizaga"
    # nomli qo'shimcha qator (nomi o'zgarishi mumkin) - loyiha egasi
    # tasdiqlagan: bu doim dollarda, naqd_dollar'ga qo'shiladi.
    wb = Workbook()
    ws = wb.active
    ws.append([None, "Rashid aka", "Jasur aka"])
    ws.append(["naqt", 28500000, None])
    ws.append(["Qog'oz", 7400, None])
    ws.append(["Plastik", None, None])
    ws.append(["click", 2500000, None])
    ws.append(["Vizaga", 1245.5, " "])  # bo'sh joy - raqam emas, e'tiborsiz
    path = tmp_path / "hisobot_vizaga.xlsx"
    wb.save(path)

    result = parse_hisobot_file(path)

    assert result["Rashid aka"] == {
        "naqd_som": 28500000,
        "naqd_dollar": 8645.5,  # 7400 + 1245.5
        "terminal": 0,
        "click": 2500000,
    }
    assert result["Jasur aka"] == {
        "naqd_som": 0,
        "naqd_dollar": 0,
        "terminal": 0,
        "click": 0,
    }


def test_parse_hisobot_file_defaults_missing_row_to_zero_for_all_fields(tmp_path):
    # Fayldan "Plastik" qatori butunlay tushib qolgan kun - natija baribir
    # barcha 4 maydonni o'z ichiga olishi kerak (terminal=0), report_parser
    # bilan bir xil kafolat: DailyEntry/diff mantig'i har doim to'liq
    # {naqd_som, click, naqd_dollar, terminal} dict kutadi.
    wb = Workbook()
    ws = wb.active
    ws.append([None, "Rashid aka"])
    ws.append(["naqt", 100000])
    ws.append(["Qog'oz", 20])
    ws.append(["click", 5000])
    path = tmp_path / "hisobot_no_plastik.xlsx"
    wb.save(path)

    result = parse_hisobot_file(path)

    assert result["Rashid aka"] == {
        "naqd_som": 100000,
        "naqd_dollar": 20,
        "terminal": 0,
        "click": 5000,
    }
