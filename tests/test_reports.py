from datetime import date

from openpyxl import load_workbook

from reports import build_debtors_report


def test_build_debtors_report_writes_sorted_rows_with_total(tmp_path):
    debtors = [
        ("Alisher aka", 950.5, date(2026, 8, 1), 50.0),
        ("Bobur aka", 200.0, None, 0),
    ]
    daily_rows = []
    path = tmp_path / "qarzdorlar.xlsx"

    build_debtors_report(debtors, daily_rows, path, sana=date(2026, 8, 18))

    wb = load_workbook(path)
    ws = wb["Qarzdorlar"]

    assert ws["A1"].value == "Qarzdorlar ro'yxati - 2026-08-18"
    assert ws["A2"].value == "Jami qarzdorlar: 2 ta"

    header_row = 4
    assert [ws.cell(row=header_row, column=c).value for c in range(1, 6)] == [
        "№", "Kontragent", "Oxirgi to'lov sanasi", "Oxirgi to'lov summasi ($)", "Qarz ($)",
    ]

    assert ws.cell(row=5, column=2).value == "Alisher aka"
    assert ws.cell(row=5, column=3).value == "2026-08-01"
    assert ws.cell(row=5, column=4).value == 50.0
    assert ws.cell(row=5, column=5).value == 950.5

    assert ws.cell(row=6, column=2).value == "Bobur aka"
    assert ws.cell(row=6, column=3).value == "ma'lum emas"
    assert ws.cell(row=6, column=4).value == 0
    assert ws.cell(row=6, column=5).value == 200.0

    assert ws.cell(row=7, column=2).value == "JAMI"
    assert ws.cell(row=7, column=5).value == 1150.5


def test_build_debtors_report_empty_debtors(tmp_path):
    path = tmp_path / "qarzdorlar_bosh.xlsx"

    build_debtors_report([], [], path, sana=date(2026, 8, 18))

    wb = load_workbook(path)
    ws = wb["Qarzdorlar"]
    assert ws["A2"].value == "Jami qarzdorlar: 0 ta"
    assert ws.cell(row=5, column=2).value == "JAMI"
    assert ws.cell(row=5, column=5).value == 0


def test_build_debtors_report_daily_payments_sheet(tmp_path):
    daily_rows = [
        (date(2026, 8, 2), 64000, 0, 0, 0, 5.0),
        (date(2026, 8, 1), 128000, 12800, 0, 5, 16.0),
    ]
    path = tmp_path / "qarzdorlar_kunlik.xlsx"

    build_debtors_report([], daily_rows, path, sana=date(2026, 8, 18))

    wb = load_workbook(path)
    assert "Kunlik to'lovlar" in wb.sheetnames
    ws = wb["Kunlik to'lovlar"]

    header_row = 3
    assert [ws.cell(row=header_row, column=c).value for c in range(1, 7)] == [
        "Sana", "Naqd so'm", "Click", "Terminal", "Naqd dollar", "Jami ($)",
    ]

    assert ws.cell(row=4, column=1).value == "2026-08-02"
    assert ws.cell(row=4, column=2).value == 64000
    assert ws.cell(row=4, column=6).value == 5.0

    assert ws.cell(row=5, column=1).value == "2026-08-01"
    assert ws.cell(row=5, column=2).value == 128000
    assert ws.cell(row=5, column=3).value == 12800
    assert ws.cell(row=5, column=6).value == 16.0

    total_row = 6
    assert ws.cell(row=total_row, column=1).value == "JAMI"
    assert ws.cell(row=total_row, column=2).value == 192000
    assert ws.cell(row=total_row, column=3).value == 12800
    assert ws.cell(row=total_row, column=6).value == 21.0
